# -*- coding: utf-8 -*-
"""
generate_report_docx.py
生成 AdaptSort 算法设计报告（DOCX格式）
严格保持 adaptsort.docx 与 一些细节讲解.docx 的亲切、透彻、高度专业与工程严密的语言风格。
"""

import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, fill_hex):
    """设置单元格背景颜色"""
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    """设置单元格内边距 (单位: dxa, 1 pt = 20 dxa)"""
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(
        f'<w:tcMar {nsdecls("w")}>'
        f'<w:top w:w="{top}" w:type="dxa"/>'
        f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
        f'<w:left w:w="{left}" w:type="dxa"/>'
        f'<w:right w:w="{right}" w:type="dxa"/>'
        f'</w:tcMar>'
    )
    tcPr.append(tcMar)

def set_table_borders(table, color="D0D7DE", sz="4", val="single"):
    """设置表格细灰色边框"""
    tblPr = table._tbl.tblPr
    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        f'<w:top w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:bottom w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:left w:val="none"/>'
        f'<w:right w:val="none"/>'
        f'<w:insideH w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:insideV w:val="none"/>'
        f'</w:tblBorders>'
    )
    tblPr.append(borders)

def format_run(run, font_name="宋体", size_pt=10.5, bold=False, italic=False, color_rgb=(51, 51, 51)):
    """统一设置西文与中文字体"""
    run.font.name = font_name
    run.font.size = Pt(size_pt)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = RGBColor(*color_rgb)
    rPr = run._r.get_or_add_rPr()
    rFonts = parse_xml(
        f'<w:rFonts {nsdecls("w")} '
        f'w:ascii="{font_name}" '
        f'w:hAnsi="{font_name}" '
        f'w:eastAsia="{font_name}" '
        f'w:cs="{font_name}"/>'
    )
    rPr.append(rFonts)

def add_paragraph_styled(doc, text="", font_name="宋体", size_pt=10.5, bold=False, italic=False,
                         color_rgb=(51, 51, 51), space_before=0, space_after=6, line_spacing=1.3,
                         align=WD_ALIGN_PARAGRAPH.LEFT):
    """添加带有指定字体和段落间距的段落"""
    p = doc.add_paragraph()
    p.alignment = align
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = line_spacing
    if text:
        run = p.add_run(text)
        format_run(run, font_name=font_name, size_pt=size_pt, bold=bold, italic=italic, color_rgb=color_rgb)
    return p

def add_heading_1(doc, text):
    """添加一级标题"""
    return add_paragraph_styled(
        doc, text, font_name="黑体", size_pt=15.0, bold=True, color_rgb=(15, 23, 42),
        space_before=16, space_after=8, line_spacing=1.2
    )

def add_heading_2(doc, text):
    """添加二级标题"""
    return add_paragraph_styled(
        doc, text, font_name="黑体", size_pt=12.5, bold=True, color_rgb=(30, 41, 59),
        space_before=12, space_after=6, line_spacing=1.2
    )

def add_heading_3(doc, text):
    """添加三级标题"""
    return add_paragraph_styled(
        doc, text, font_name="黑体", size_pt=11.0, bold=True, color_rgb=(51, 65, 85),
        space_before=8, space_after=4, line_spacing=1.2
    )

def add_code_block(doc, code_text):
    """添加高可读的代码块"""
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.cell(0, 0)
    set_cell_background(cell, "F8FAFC")
    set_cell_margins(cell, top=140, bottom=140, left=200, right=200)

    # 左边框浅蓝加粗，其余无边框
    tcPr = cell._tc.get_or_add_tcPr()
    borders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        f'<w:top w:val="none"/>'
        f'<w:left w:val="single" w:sz="24" w:space="0" w:color="3B82F6"/>'
        f'<w:bottom w:val="none"/>'
        f'<w:right w:val="none"/>'
        f'</w:tcBorders>'
    )
    tcPr.append(borders)

    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.15
    run = p.add_run(code_text)
    format_run(run, font_name="Consolas", size_pt=9.0, bold=False, color_rgb=(30, 41, 59))
    
    # 后面空一行
    add_paragraph_styled(doc, "", size_pt=4, space_after=4)

def add_callout(doc, text, title="关键细节提示：", fill_hex="EFF6FF", border_hex="3B82F6"):
    """添加高亮提示框"""
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.cell(0, 0)
    set_cell_background(cell, fill_hex)
    set_cell_margins(cell, top=120, bottom=120, left=180, right=180)

    tcPr = cell._tc.get_or_add_tcPr()
    borders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        f'<w:top w:val="none"/>'
        f'<w:left w:val="single" w:sz="20" w:space="0" w:color="{border_hex}"/>'
        f'<w:bottom w:val="none"/>'
        f'<w:right w:val="none"/>'
        f'</w:tcBorders>'
    )
    tcPr.append(borders)

    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.25

    if title:
        run_t = p.add_run(title + " ")
        format_run(run_t, font_name="黑体", size_pt=9.5, bold=True, color_rgb=(30, 58, 138))

    run = p.add_run(text)
    format_run(run, font_name="宋体", size_pt=9.5, bold=False, color_rgb=(30, 41, 59))
    add_paragraph_styled(doc, "", size_pt=4, space_after=4)

def create_table_styled(doc, headers, data, col_widths=None):
    """创建美化表格"""
    table = doc.add_table(rows=len(data) + 1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(table)

    # 填充表头
    for col_idx, h in enumerate(headers):
        cell = table.cell(0, col_idx)
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        set_cell_background(cell, "F1F5F9")
        set_cell_margins(cell, top=140, bottom=140, left=140, right=140)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(2)
        run = p.add_run(h)
        format_run(run, font_name="黑体", size_pt=9.5, bold=True, color_rgb=(15, 23, 42))

    # 填充数据行
    for row_idx, row_data in enumerate(data):
        bg = "FFFFFF" if row_idx % 2 == 0 else "F8FAFC"
        for col_idx, val in enumerate(row_data):
            cell = table.cell(row_idx + 1, col_idx)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            set_cell_background(cell, bg)
            set_cell_margins(cell, top=100, bottom=100, left=140, right=140)
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.line_spacing = 1.2
            run = p.add_run(str(val))
            format_run(run, font_name="宋体", size_pt=9.0, bold=False, color_rgb=(51, 65, 85))

    # 设置列宽
    if col_widths:
        for row in table.rows:
            for idx, w in enumerate(col_widths):
                row.cells[idx].width = Inches(w)

    add_paragraph_styled(doc, "", size_pt=4, space_after=6)
    return table

print("Helper functions initialized.")
